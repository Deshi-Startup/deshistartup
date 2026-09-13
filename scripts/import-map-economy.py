"""Import checked Economic Census 2024 totals and permanent-sector tables.

Default operation needs only Python's standard library and checked source CSVs.
--extract-from DIR recreates CSVs from the hashed official XLSX and pdftotext
inputs in DIR; this extraction mode additionally requires openpyxl.
"""
import argparse
import csv
import hashlib
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data/maps/sources/economic-2024"
OUTPUT = ROOT / "data/maps/economy.json"
NATIONAL = "country-bangladesh"
CODES = list("BCDEFGHIJKLMNOPQRS")
LOCALITIES = {"Total", "Rural", "Urban"}
METRICS = ["economicUnits", "permanentEstablishments", "temporaryEstablishments",
           "economicHouseholds", "personsEngaged", "permanentPersonsEngaged",
           "temporaryPersonsEngaged", "economicHouseholdPersonsEngaged"]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def name_key(name):
    return re.sub(r"[^a-z]", "", name.lower()).replace("chapainababganj", "chapainawabganj")


def geography():
    rows = json.loads((ROOT / "data/maps/regions.json").read_text(encoding="utf-8"))["regions"]
    require(len(rows) == 72 and len({r["id"] for r in rows}) == 72, "Expected 72 unique canonical regions")
    lookup = {name_key(r["name"]["en"] + (" Division" if r["level"] == "division" else "")): r["id"] for r in rows}
    lookup["national"] = NATIONAL
    return rows, lookup


def write_csv(path, fields, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def extract(raw_dir, source_dir=SOURCE):
    import openpyxl
    metadata = json.loads((source_dir / "metadata.json").read_text(encoding="utf-8"))
    for filename, expected in metadata["source"]["inputFiles"].items():
        path = raw_dir / filename
        require(path.stat().st_size == expected["bytes"] and digest(path) == expected["sha256"],
                f"Source bytes/hash differ from reviewed release: {filename}")
    canonical, lookup = geography()
    detail = openpyxl.load_workbook(raw_dir / "economic-detail.xlsx", read_only=True, data_only=True)
    overview = openpyxl.load_workbook(raw_dir / "economic-overview.xlsx", read_only=True, data_only=True)

    # Match the complete eight-count row against the final PDF, independently of XLSX.
    v1pages = (raw_dir / "economic-final.txt").read_text(encoding="utf-8").split("\f")
    pdf_rows = {}
    for page in range(143, 148):
        for line in v1pages[page - 1].splitlines():
            numbers = tuple(int(n) for n in re.findall(r"(?<!\S)\d+(?!\S)", line))
            if len(numbers) == 8:
                require(numbers not in pdf_rows, "Ambiguous duplicate S2 numeric row in PDF")
                pdf_rows[numbers] = page
    totals = []
    region_id = None
    source_name = None
    for row_number, row in enumerate(detail["Table S2"].iter_rows(min_row=5, values_only=True), 5):
        name = str(row[0]).strip()
        locality = name if name in {"Rural", "Urban"} else "Total"
        if locality == "Total":
            require(name_key(name) in lookup, f"Unknown S2 geography: {name}")
            region_id, source_name = lookup[name_key(name)], name
        values = tuple(row[1:9])
        require(all(type(v) is int and v >= 0 for v in values), f"Non-integer S2 counts at row {row_number}")
        require(values in pdf_rows, f"S2 XLSX row {row_number} has no identical eight-count PDF row")
        totals.append({"id": region_id, "sourceName": source_name, "locality": locality,
                       **dict(zip(METRICS, values)), "xlsxRow": row_number,
                       "pdfPage": pdf_rows[values], "printedPage": pdf_rows[values] - 24})

    checks = []
    division_overview = {}
    for row_number, row in enumerate(overview["Table 3.3"].iter_rows(min_row=7, values_only=True), 7):
        name = str(row[0]).strip()
        if name in {"Rural", "Urban"}:
            continue
        region_id = NATIONAL if name == "National" else lookup[name_key(name + " Division")]
        division_overview[region_id] = ([row[i] for i in [1, 3, 5, 7]], row_number)
    for row_number, row in enumerate(detail["Table S3"].iter_rows(min_row=6, values_only=True), 6):
        region_id = lookup[name_key(row[0])]
        require(sum(row[i] for i in [3, 5, 7, 9, 11]) == row[1] and
                sum(row[i] for i in [4, 6, 8, 10, 12]) == row[2], f"S3 category totals differ: {region_id}")
        require(row[13] == row[7] + row[9] and row[14] == row[8] + row[10], f"S3 SME overlap mismatch: {region_id}")
        o = division_overview.get(region_id)
        checks.append({"id": region_id, "s3EconomicUnits": row[1], "s3PersonsEngaged": row[2], "s3XlsxRow": row_number,
                       "overviewEconomicUnits": o[0][0] if o else "", "overviewPermanentEstablishments": o[0][1] if o else "",
                       "overviewTemporaryEstablishments": o[0][2] if o else "", "overviewEconomicHouseholds": o[0][3] if o else "",
                       "overviewXlsxRow": o[1] if o else ""})

    v2pages = (raw_dir / "economic-vol2.txt").read_text(encoding="utf-8").split("\f")
    sector_tables = {}
    for table, first, last in [("S4", 58, 67), ("S5", 67, 77)]:
        rows = []
        region_id = None
        source_name = None
        for page_number in range(first, last + 1):
            page = v2pages[page_number - 1]
            # Both tables share physical page 67; never parse one as the other.
            if page_number == 67:
                page = page.split("Table S5:")[0 if table == "S4" else 1]
            lines = page.splitlines()
            for index, line in enumerate(lines):
                match = re.match(r"^\s*(.*?)\b(Total|Rural|Urban)\s+([\d\s]+)$", line)
                if not match:
                    continue
                name, locality = match[1].strip(), match[2]
                if name and index + 1 < len(lines) and lines[index + 1].strip() == "Division":
                    name += " Division"
                if locality == "Total":
                    require(name_key(name) in lookup, f"Unknown Volume II {table} geography: {name}")
                    region_id, source_name = lookup[name_key(name)], name
                require(region_id is not None, f"Missing geography before {table} locality row")
                counts = list(map(int, match[3].split()))
                require(len(counts) == 19, f"Expected all-sectors total plus 18 sectors: {table} {page_number} {name}")
                rows.append({"id": region_id, "sourceName": source_name, "locality": locality,
                             "allSectors": counts[0], **dict(zip(CODES, counts[1:])),
                             "pdfPage": page_number, "printedPage": page_number - 14})
        sector_tables[table] = rows

    # Validate all extracted tables before changing source CSVs.
    validate(totals, checks, sector_tables["S4"], sector_tables["S5"], canonical)
    write_csv(source_dir / "totals.csv", list(totals[0]), totals)
    write_csv(source_dir / "regional-checks.csv", list(checks[0]), checks)
    write_csv(source_dir / "permanent-sector-establishments.csv", list(sector_tables["S4"][0]), sector_tables["S4"])
    write_csv(source_dir / "permanent-sector-persons.csv", list(sector_tables["S5"][0]), sector_tables["S5"])


def read_csv(path):
    with path.open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for row in rows:
        for key, value in row.items():
            if key not in {"id", "sourceName", "locality"} and value != "":
                require(re.fullmatch(r"\d+", value) is not None, f"Non-integer source count: {path.name} {key}")
                row[key] = int(value)
    return rows


def index_table(rows, ids, label):
    indexed = {(r["id"], r["locality"]): r for r in rows}
    require(len(indexed) == len(rows) == 219, f"{label}: expected 219 unique geography/locality rows")
    require(set(indexed) == {(i, l) for i in ids for l in LOCALITIES}, f"{label}: geography/locality coverage differs")
    return indexed


def validate(totals, checks, establishments, persons, canonical):
    ids = {r["id"] for r in canonical} | {NATIONAL}
    t = index_table(totals, ids, "S2")
    e = index_table(establishments, ids, "S4")
    p = index_table(persons, ids, "S5")
    c = {r["id"]: r for r in checks}
    require(len(c) == len(checks) == 73 and set(c) == ids, "S3 cross-check coverage differs")
    _, names = geography()
    for key, row in t.items():
        require(names.get(name_key(row["sourceName"])) == row["id"], f"S2 source-name/id mismatch: {key}")
        require(sum(row[k] for k in METRICS[1:4]) == row["economicUnits"], f"Unit-type sum differs: {key}")
        require(sum(row[k] for k in METRICS[5:8]) == row["personsEngaged"], f"TPE-type sum differs: {key}")
        require(143 <= row["pdfPage"] <= 147 and row["printedPage"] == row["pdfPage"] - 24, f"Unexpected S2 page: {key}")
        for table, metric, first, last in [(e, "permanentEstablishments", 58, 67), (p, "permanentPersonsEngaged", 67, 77)]:
            r = table[key]
            require(names.get(name_key(r["sourceName"])) == r["id"], f"Sector source-name/id mismatch: {key}")
            require(first <= r["pdfPage"] <= last and r["printedPage"] == r["pdfPage"] - 14, f"Unexpected sector page: {key}")
            require(sum(r[code] for code in CODES) == r["allSectors"] == row[metric], f"Sector totals differ from S2: {key} {metric}")
        if key[1] == "Total":
            check = c[key[0]]
            require((row["economicUnits"], row["personsEngaged"]) == (check["s3EconomicUnits"], check["s3PersonsEngaged"]), f"S2/S3 mismatch: {key}")
            has_overview = key[0] == NATIONAL or key[0].startswith("division-")
            require(bool(check["overviewXlsxRow"]) == has_overview, f"Overview coverage mismatch: {key}")
            if has_overview:
                require([row[k] for k in METRICS[:4]] == [check[k] for k in ["overviewEconomicUnits", "overviewPermanentEstablishments", "overviewTemporaryEstablishments", "overviewEconomicHouseholds"]], f"S2/3.3 mismatch: {key}")

    # Check locality partitions, every sector, division sums and national sums.
    for table, metrics in [(t, METRICS), (e, ["allSectors", *CODES]), (p, ["allSectors", *CODES])]:
        for region_id in ids:
            for metric in metrics:
                require(table[region_id, "Total"][metric] == sum(table[region_id, l][metric] for l in ["Rural", "Urban"]), f"Rural/urban mismatch: {region_id} {metric}")
        for locality in LOCALITIES:
            for metric in metrics:
                for division in [r for r in canonical if r["level"] == "division"]:
                    districts = [r["id"] for r in canonical if r["level"] == "district" and r["division"] == division["key"]]
                    require(sum(table[d, locality][metric] for d in districts) == table[division["id"], locality][metric], f"District/division mismatch: {division['id']} {locality} {metric}")
                for level in ["district", "division"]:
                    require(sum(table[r["id"], locality][metric] for r in canonical if r["level"] == level) == table[NATIONAL, locality][metric], f"National/{level} mismatch: {locality} {metric}")
    return t, c, e, p


def build(source_dir=SOURCE):
    result = json.loads((source_dir / "metadata.json").read_text(encoding="utf-8"))
    require(list(result["sectors"]) == CODES, "Expected published BSIC B–S sector order")
    require(result["source"]["sectorShareDenominator"] == "permanentEstablishments", "Sector shares require permanent-establishment denominator")
    filenames = ["totals.csv", "regional-checks.csv", "permanent-sector-establishments.csv", "permanent-sector-persons.csv"]
    canonical, _ = geography()
    t, c, e, p = validate(*(read_csv(source_dir / f) for f in filenames), canonical)

    def observation(region_id):
        key = (region_id, "Total")
        r = t[key]
        row = {metric: r[metric] for metric in METRICS}
        row["source"] = {"volume": "I", "table": "S2", "sheet": "Table S2", "xlsxRow": r["xlsxRow"], "pdfPage": r["pdfPage"], "printedPage": r["printedPage"], "sourceName": r["sourceName"]}
        row["crosscheckSource"] = {"volume": "I", "table": "S3", "sheet": "Table S3", "xlsxRow": c[region_id]["s3XlsxRow"]}
        row["sectors"] = {code: {"permanentEstablishments": e[key][code], "personsEngaged": p[key][code]} for code in CODES}
        for field, table, number in [("sectorEstablishmentsSource", e, "S4"), ("sectorPersonsSource", p, "S5")]:
            row[field] = {"volume": "II", "table": number, "pdfPage": table[key]["pdfPage"], "printedPage": table[key]["printedPage"], "sourceName": table[key]["sourceName"]}
        return row

    result["national"] = observation(NATIONAL)
    result["regions"] = {r["id"]: observation(r["id"]) for r in canonical}
    result["validation"].update({"districts": 64, "divisions": 8, "sectors": 18,
                                  "localityRowsPerTable": 219, "unresolvedCells": 0,
                                  "inputSHA256": {f: digest(source_dir / f) for f in ["metadata.json", *filenames]}})
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=SOURCE)
    parser.add_argument("--extract-from", type=Path, help="Directory containing the six hashed PDF/XLSX/text inputs")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    require(not (args.check and args.extract_from), "--check does not write source transcriptions; run --extract-from separately")
    if args.extract_from:
        extract(args.extract_from, args.source_dir)
    text = json.dumps(build(args.source_dir), ensure_ascii=False, indent=2) + "\n"
    if args.check:
        require(args.output.read_text(encoding="utf-8") == text, "economy.json differs from checked sources; rerun importer")
    else:
        args.output.write_text(text, encoding="utf-8")
    print("Validated 64 districts, eight divisions, national total and all 18 permanent-establishment sectors; all row, locality and geographic totals reconcile.")


if __name__ == "__main__":
    main()
